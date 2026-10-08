import streamlit as st
import pandas as pd
import openpyxl
import os
from datetime import date
from dotenv import load_dotenv

import auth

auth.exigir_login(
    painel_nome="Atualização Fornecedor",
    titulo_painel="Atualização de Fornecedores",
    subtitulo="Relação de Fornecedores Atualizados",
    icone="🏢"
)

import sys
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..')))
from onedrive_downloader import download_excel_from_onedrive, upload_excel_to_onedrive

ONEDRIVE_PATH = 'Administrativo/Qualidade/Área de Cadastros/Painel Gestão de Cadastros/Controle Cadastros.xlsx'
FILE_PATH = 'Administrativo/Qualidade/Área de Cadastros/Painel Gestão de Cadastros/Controle Cadastros.xlsx' # Mantido apenas se houver gravação posterior
NOME_ABA = 'Alt_Att Fornec'

def formata_zerados():
    if st.session_state.codigo_input:
        st.session_state.codigo_input = st.session_state.codigo_input.zfill(6)
    if st.session_state.loja_input:
        st.session_state.loja_input = st.session_state.loja_input.zfill(2)

def limpar_tudo():
    for key in list(st.session_state.keys()):
        if key == 'atualizacao_input':
            st.session_state[key] = None
        elif key == 'check_branco':
            st.session_state[key] = False
        # Não apaga a data, a flag de limpar, a busca, nem a flag de sucesso
        elif key in ['data_input', 'deve_limpar', 'busca_fornecedor', 'mostrar_sucesso', 'autenticado_forn', 'usuario_login_forn', 'senha_login_forn', 'btn_sair_forn', 'is_root', 'user_info', 'username_logado', 'usuario_primeiro_login'] or key.startswith('auth_'):
            pass 
        else:
            st.session_state[key] = ""

# Flag para limpar cache sem bugar a tela
if st.session_state.get('deve_limpar', False):
    limpar_tudo()
    st.session_state['deve_limpar'] = False

# --- FUNÇÃO DO POP-UP (MODAL) ---
@st.dialog("📝 Cadastrar Nova Atualização", width='medium')
def modal_novo_registro():
    c1, c2, c3, c4 = st.columns([1, 0.8, 1.5, 2])
    codigo = c1.text_input("Código", max_chars=6, key="codigo_input", on_change=formata_zerados)
    loja = c2.text_input("Loja", max_chars=2, key="loja_input", on_change=formata_zerados)
    
    data_solic = c3.date_input("Data Solicitação", date.today(), format="DD/MM/YYYY", key="data_input")
    
    atualizacao = c4.selectbox("O que será atualizado?", 
                             ["Dados Bancários", "Dados Cadastrais"], 
                             index=None, 
                             placeholder="Selecione uma opção...",
                             key="atualizacao_input")

    c5, c6 = st.columns([3, 1])
    razao_social = c5.text_input("Razão Social", key="razao_input")
    solicitante = c6.text_input("Solicitante", key="solicitante_input")

    if atualizacao:
        st.write("---")
        
        dados_em_branco = st.checkbox("Dados em branco", key="check_branco")

        dados_antigos_val = ""
        dados_novos_val = ""

        if atualizacao == "Dados Bancários":
            st.markdown("**Dados Bancários ANTIGOS:**")
            if dados_em_branco:
                st.info("Campo preenchido automaticamente como 'Em branco'.")
                dados_antigos_val = "Em branco"
            else:
                ant_col1, ant_col2, ant_col3 = st.columns(3)
                b_ant = ant_col1.text_input("Banco", key="b_ant")
                a_ant = ant_col2.text_input("Agência", key="a_ant")
                c_ant = ant_col3.text_input("Conta", key="c_ant")
                dados_antigos_val = f"Banco: {b_ant} - Ag: {a_ant} - C/C: {c_ant}"
            st.write("---")

            st.markdown("**Dados Bancários NOVOS:**")
            nov_col1, nov_col2, nov_col3 = st.columns(3)
            b_nov = nov_col1.text_input("Banco", key="b_nov")
            a_nov = nov_col2.text_input("Agência", key="a_nov")
            c_nov = nov_col3.text_input("Conta", key="c_nov")
            dados_novos_val = f"Banco: {b_nov} - Ag: {a_nov} - C/C: {c_nov}"
            st.write("---")

        elif atualizacao == "Dados Cadastrais":
            st.markdown("**Dados Cadastrais:**")
            if dados_em_branco:
                st.info("Campo preenchido automaticamente como 'Em branco'.")
                dados_antigos_val = "Em branco"
            else:
                dados_antigos_val = st.text_area("Dados Antigos (Cadastrais)", height=100, key="cad_antigo")
                
            dados_novos_val = st.text_area("Dados Novos (Cadastrais)", height=100, key="cad_novo")

        obs = st.text_input("Observação", key="obs_input")

        st.write("---")
        
        # --- NOVOS CHECKBOXES ---
        col_chk1, col_chk2 = st.columns(2)
        with col_chk1:
            auth_pagamento = st.checkbox("Possui autorização de pagamento?", key="chk_auth")
        with col_chk2:
            fornecedor_novo = st.checkbox("É um fornecedor novo?", key="chk_novo")

        st.write("---")

        col_btn1, col_btn2 = st.columns(2)
        with col_btn1:
            submit = st.button("💾 Salvar", use_container_width=True, type="primary")
        with col_btn2:
            if st.button("🧹 Limpar Campos", use_container_width=True):
                st.session_state['deve_limpar'] = True
                st.rerun()

        if submit:
            # --- VALIDAÇÃO DE REGRAS DE NEGÓCIO ---
            if auth_pagamento and not obs.strip():
                st.warning("Como possui autorização de pagamento, o campo 'Observação' é obrigatório!")
            else:
                codigo_final = st.session_state.codigo_input
                loja_final = st.session_state.loja_input

                if not codigo_final or not loja_final or not razao_social:
                    st.warning("Preencha o Código, a Loja e a Razão Social para registrar na planilha!")
                else:
                    try:
                        # 1. Baixar o arquivo original da nuvem para a memória
                        file_bytes_io = download_excel_from_onedrive(ONEDRIVE_PATH)
                        wb = openpyxl.load_workbook(file_bytes_io)
                        ws = wb[NOME_ABA]

                        next_row = 1
                        for r in range(ws.max_row, 0, -1):
                            if ws.cell(row=r, column=1).value is not None:
                                next_row = r + 1
                                break

                        # Tratamento de Sim/Não para as novas colunas
                        val_auth = "Sim" if auth_pagamento else "Não"
                        val_novo = "Sim" if fornecedor_novo else "Não"

                        ws.cell(row=next_row, column=1, value=codigo_final)
                        ws.cell(row=next_row, column=2, value=loja_final)
                        ws.cell(row=next_row, column=3, value=razao_social)
                        ws.cell(row=next_row, column=4, value=data_solic.strftime('%d/%m/%Y'))
                        ws.cell(row=next_row, column=5, value=atualizacao)
                        ws.cell(row=next_row, column=6, value=dados_antigos_val)
                        ws.cell(row=next_row, column=7, value=dados_novos_val)
                        ws.cell(row=next_row, column=8, value=solicitante)
                        ws.cell(row=next_row, column=9, value=obs)
                        ws.cell(row=next_row, column=10, value=val_auth) # Coluna J
                        ws.cell(row=next_row, column=11, value=val_novo) # Coluna K

                        # 2. Salvar as alterações em memória
                        from io import BytesIO
                        out_bytes_io = BytesIO()
                        wb.save(out_bytes_io)
                        out_bytes = out_bytes_io.getvalue()
                        
                        # 3. Enviar o arquivo atualizado de volta para o SharePoint
                        upload_excel_to_onedrive(ONEDRIVE_PATH, out_bytes)
                        
                        # Limpa o cache do download para ele puxar os dados novos na proxima leitura
                        download_excel_from_onedrive.clear()
                        
                        st.session_state['mostrar_sucesso'] = True
                        st.session_state['deve_limpar'] = True
                        st.rerun()
                        
                    except PermissionError:
                        st.error("O arquivo Excel está aberto. Feche-o antes de salvar.")
                    except Exception as e:
                        st.error("Erro interno ao salvar.")

# --- TELA PRINCIPAL ---
def main():
    import sys
    sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..')))
    try:
        import ui_utils
        ui_utils.render_standard_panel(
            title="Atualização de Fornecedores",
            subtitle="Relação de Fornecedores Atualizados",
            icon_name="Fornecedores logo.png"
        )
    except ImportError:
        st.title("Atualização de Fornecedores")

    if st.session_state.get('mostrar_sucesso', False):
        st.success("Tudo certo! Registro inserido na planilha com sucesso.")
        st.session_state['mostrar_sucesso'] = False 

    st.divider()

    col_hist, col_btn = st.columns([4, 1])
    with col_hist:
        st.subheader("Histórico de Atualizações")
    with col_btn:
        if st.button("➕ Novo Registro", use_container_width=True, type="primary"):
            limpar_tudo() # <-- Puxa a descarga aqui antes de abrir a tela
            modal_novo_registro()

    termo_busca = st.text_input("🔍 Buscar por Código ou Razão Social:", key="busca_fornecedor")

    try:
        arquivo_excel = download_excel_from_onedrive(ONEDRIVE_PATH)
        df_view = pd.read_excel(arquivo_excel, sheet_name=NOME_ABA, header=1)
        df_view.columns = df_view.columns.str.strip()
        
        if 'Data Solicitação' in df_view.columns:
            df_view['Data Solicitação'] = pd.to_datetime(df_view['Data Solicitação'], errors='coerce').dt.strftime('%d/%m/%Y')
        
        if 'Código' in df_view.columns:
            df_view['Código'] = df_view['Código'].apply(lambda x: str(x).replace('.0', '').zfill(6) if pd.notnull(x) and str(x) != 'nan' else '')
        if 'Loja' in df_view.columns:
            df_view['Loja'] = df_view['Loja'].apply(lambda x: str(x).replace('.0', '').zfill(2) if pd.notnull(x) and str(x) != 'nan' else '')

        if termo_busca:
            termo = termo_busca.lower()
            df_view = df_view[
                df_view['Código'].astype(str).str.lower().str.contains(termo, na=False) |
                df_view['Razão Social'].astype(str).str.lower().str.contains(termo, na=False)
            ]

        df_view = df_view.iloc[::-1]

        st.dataframe(df_view, use_container_width=True)
    except Exception as e:
        st.error("Falha crítica ao acessar o SharePoint.")
        st.stop()

main()
