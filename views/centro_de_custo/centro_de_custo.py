import streamlit as st
import pandas as pd
import io
import os
import openpyxl
import re
import unicodedata
from datetime import datetime
import json
from dotenv import load_dotenv

load_dotenv()

import sys
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..')))

# Caminho da sua planilha gerada pelo bot
caminho_fixo = "/app/data/Aprovadores.xlsx"


# --- CSS CUSTOMIZADO ---
css = """
<style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');

    .stApp { font-family: 'Inter', sans-serif; }

    header[data-testid="stHeader"] { background-color: transparent; }

    /* === CABEÇALHO DA PÁGINA === */
    .page-header {
        position: relative;
        padding: 6px 0 20px 0;
        margin-bottom: 26px;
    }
    .page-header::after {
        content: "";
        position: absolute;
        bottom: 0; left: 0; right: 0;
        height: 2px;
        background: linear-gradient(to right, #3b82f6 0%, rgba(59, 130, 246, 0.25) 45%, transparent 100%);
        border-radius: 2px;
    }
    .page-header-content {
        display: flex;
        align-items: center;
        gap: 14px;
    }
    .page-header-icon {
        font-size: 26px;
        background: rgba(59, 130, 246, 0.1);
        border: 1px solid rgba(59, 130, 246, 0.2);
        border-radius: 12px;
        width: 52px;
        height: 52px;
        display: flex;
        align-items: center;
        justify-content: center;
        flex-shrink: 0;
    }
    .page-header-title {
        font-size: 24px;
        font-weight: 800;
        color: var(--text-color);
        letter-spacing: -0.5px;
        margin: 0 0 3px 0;
        line-height: 1.2;
    }
    .page-header-subtitle {
        font-size: 13px;
        color: var(--text-color);
        opacity: 0.45;
        font-weight: 400;
        margin: 0;
    }

    /* === CARDS DE KPI === */
    .metric-card {
        background-color: var(--secondary-background-color);
        padding: 16px 18px;
        border-radius: 10px;
        display: flex;
        justify-content: space-between;
        align-items: center;
        border: 1px solid rgba(128, 128, 128, 0.15);
        border-left: 4px solid transparent;
        box-shadow: 0 1px 4px rgba(0, 0, 0, 0.04);
        height: 90px;
        transition: transform 0.15s, box-shadow 0.15s, background-color 0.2s;
    }
    .metric-card:hover {
        transform: translateY(-2px);
        box-shadow: 0 6px 20px rgba(0, 0, 0, 0.08);
    }
    .metric-card-blue:hover  { background-color: rgba(59,  130, 246, 0.06); }
    .metric-card-green:hover { background-color: rgba(16,  185, 129, 0.06); }
    .metric-card-red:hover   { background-color: rgba(239,  68,  68, 0.06); }
    .metric-card-purple:hover{ background-color: rgba(139,  92, 246, 0.06); }

    .metric-info {
        display: flex;
        flex-direction: column;
        justify-content: center;
    }
    .metric-title {
        color: var(--text-color);
        opacity: 0.5;
        font-size: 10px;
        font-weight: 700;
        text-transform: uppercase;
        letter-spacing: 0.7px;
        margin-bottom: 6px;
        white-space: nowrap;
    }
    .metric-value {
        color: var(--text-color);
        font-size: 28px;
        font-weight: 700;
        line-height: 1.1;
    }
    .metric-chart {
        width: 45px;
        height: 45px;
        display: flex;
        align-items: center;
        justify-content: center;
    }
    .donut-green {
        width: 40px; height: 40px;
        border-radius: 50%;
        background: conic-gradient(#10b981 75%, rgba(128,128,128,0.2) 0);
        display: flex; justify-content: center; align-items: center;
    }
    .donut-green::before { content: ""; width: 24px; height: 24px; background: var(--secondary-background-color); border-radius: 50%; }
    .donut-red {
        width: 40px; height: 40px;
        border-radius: 50%;
        background: conic-gradient(#ef4444 25%, rgba(128,128,128,0.2) 0);
        display: flex; justify-content: center; align-items: center;
    }
    .donut-red::before { content: ""; width: 24px; height: 24px; background: var(--secondary-background-color); border-radius: 50%; }

    /* === TABELA === */
    .table-container {
        width: 100%;
        max-height: 500px;
        overflow-y: auto;
        overflow-x: auto;
        margin-top: 12px;
        background-color: var(--background-color);
        border-radius: 8px;
        border: 1px solid rgba(128, 128, 128, 0.15);
        box-shadow: 0 1px 6px rgba(0, 0, 0, 0.03);
    }
    .table-container::-webkit-scrollbar { width: 6px; height: 6px; }
    .table-container::-webkit-scrollbar-track { background: var(--background-color); border-radius: 4px; }
    .table-container::-webkit-scrollbar-thumb { background: rgba(128, 128, 128, 0.25); border-radius: 4px; }
    .table-container::-webkit-scrollbar-thumb:hover { background: rgba(128, 128, 128, 0.45); }

    .custom-table {
        width: 100%;
        border-collapse: collapse;
        font-size: 12px;
        color: var(--text-color);
        white-space: nowrap;
    }
    .custom-table thead th {
        position: sticky;
        top: 0;
        z-index: 10;
        color: var(--text-color);
        opacity: 0.65;
        text-transform: uppercase;
        padding: 10px 12px;
        text-align: left;
        font-size: 10px;
        font-weight: 700;
        letter-spacing: 0.6px;
        border-bottom: 2px solid rgba(128, 128, 128, 0.15);
    }
    @media (prefers-color-scheme: dark) {
        .custom-table thead th { background-color: #1a1e26; }
    }
    @media (prefers-color-scheme: light) {
        .custom-table thead th { background-color: #f8fafc; }
    }
    .custom-table td {
        padding: 9px 12px;
        text-align: left;
        border-bottom: 1px solid rgba(128, 128, 128, 0.08);
    }
    .custom-table tbody tr { transition: background-color 0.15s; }
    .table-sort-link { color: inherit; text-decoration: none; cursor: pointer; display: flex; align-items: center; gap: 4px; }
    .table-sort-link:hover { color: #3b82f6; }
    .custom-table tbody tr:hover { background-color: rgba(128, 128, 128, 0.06); }
    .tr-bloqueado { background-color: rgba(245, 158, 11, 0.12) !important; }

    .empresa-col { display: flex; align-items: center; gap: 8px; font-weight: 500; }
    .action-icons { display: flex; gap: 12px; font-size: 15px; color: var(--text-color); opacity: 0.6; }
    .action-icons span { cursor: pointer; transition: all 0.15s; }
    .action-icons span:hover { color: var(--primary-color); opacity: 1; transform: scale(1.1); }
    .action-icons .delete:hover { color: #ef4444; opacity: 1; }

    /* === TAGS DE STATUS === */
    .status-tag {
        padding: 3px 10px;
        border-radius: 20px;
        font-size: 10px;
        font-weight: 700;
        display: inline-flex;
        align-items: center;
        gap: 5px;
        border: 1px solid transparent;
        text-transform: uppercase;
        letter-spacing: 0.4px;
        white-space: nowrap;
    }
    .status-ativo {
        background-color: rgba(16, 185, 129, 0.1);
        color: #10b981;
        border-color: rgba(16, 185, 129, 0.25);
    }
    .status-inativo {
        background-color: rgba(239, 68, 68, 0.1);
        color: #ef4444;
        border-color: rgba(239, 68, 68, 0.25);
    }
    .status-aviso {
        background-color: rgba(245, 158, 11, 0.1);
        color: #f59e0b;
        border-color: rgba(245, 158, 11, 0.25);
    }

    /* === DICA DA TABELA === */
    .table-hint {
        display: inline-flex;
        align-items: center;
        gap: 8px;
        padding: 7px 14px;
        background-color: rgba(59, 130, 246, 0.06);
        border: 1px solid rgba(59, 130, 246, 0.2);
        border-radius: 8px;
        font-size: 12px;
        color: #3b82f6;
        font-weight: 500;
        margin-bottom: 6px;
    }

    /* === BOTÕES === */
    .btn-limpar button {
        background-color: transparent !important;
        border: 1px solid rgba(239, 68, 68, 0.55) !important;
        color: #ef4444 !important;
        padding: 4px 15px !important;
        border-radius: 6px !important;
        height: 38px !important;
        margin-top: 28px !important;
        font-weight: 600 !important;
        font-size: 13px !important;
    }
    .btn-limpar button:hover {
        background-color: rgba(239, 68, 68, 0.08) !important;
        border-color: #ef4444 !important;
    }
    .btn-excel button {
        background-color: rgba(16, 185, 129, 0.08) !important;
        color: #10b981 !important;
        border: 1px solid rgba(16, 185, 129, 0.4) !important;
        border-radius: 6px !important;
        font-weight: 600 !important;
        font-size: 13px !important;
    }
    .btn-excel button:hover {
        background-color: rgba(16, 185, 129, 0.16) !important;
        border-color: #10b981 !important;
    }

    /* === RODAPÉ === */
    .pag-text {
        color: var(--text-color);
        opacity: 0.55;
        font-size: 13px;
        font-weight: 500;
        margin-top: 16px;
    }
</style>
"""
st.markdown(css, unsafe_allow_html=True)

if os.path.exists(caminho_fixo):
    arquivo_unico = caminho_fixo
else:
    st.sidebar.error("❌ Arquivo não encontrado no caminho padrão.")
    arquivo_unico = st.sidebar.file_uploader("Upload da Base Unificada", type=["xlsx"])


@st.cache_data(ttl=14400)
def carregar_dados(caminho_arquivo, timestamp):
    plan2 = pd.read_excel(caminho_arquivo, sheet_name='Plan2', dtype=str)
    base = pd.read_excel(caminho_arquivo, sheet_name='Base', dtype=str)
    
    plan2.columns = plan2.columns.str.strip()
    base.columns = base.columns.str.strip()
    
    plan2 = plan2.apply(lambda col: col.str.strip() if col.dtype == 'object' else col)
    base = base.apply(lambda col: col.str.strip() if col.dtype == 'object' else col)
    
    try:
        df_bloqueios = pd.read_excel(caminho_arquivo, sheet_name='Bloqueios', dtype=str)
        df_bloqueios.columns = df_bloqueios.columns.str.strip()
        if 'CTT_EMPRESA_BLOQ' not in df_bloqueios.columns:
            df_bloqueios = pd.DataFrame(columns=['CTT_EMPRESA_BLOQ', 'CTT_FILIAL_BLOQ', 'CTT_CUSTO_BLOQ', 'DATA_BLOQUEIO'])
        else:
            df_bloqueios = df_bloqueios.apply(lambda col: col.str.strip() if col.dtype == 'object' else col)
            df_bloqueios['CTT_EMPRESA_BLOQ'] = df_bloqueios['CTT_EMPRESA_BLOQ'].str.replace(r'\.0$', '', regex=True).str.zfill(2)
            df_bloqueios['CTT_FILIAL_BLOQ'] = df_bloqueios['CTT_FILIAL_BLOQ'].str.replace(r'\.0$', '', regex=True).str.zfill(6)
            df_bloqueios['CTT_CUSTO_BLOQ'] = df_bloqueios['CTT_CUSTO_BLOQ'].str.strip()
    except Exception:
        df_bloqueios = pd.DataFrame(columns=['CTT_EMPRESA_BLOQ', 'CTT_FILIAL_BLOQ', 'CTT_CUSTO_BLOQ', 'DATA_BLOQUEIO'])
        
    try:
        df_pendencias = pd.read_excel(caminho_arquivo, sheet_name='Pendencias', dtype=str)
        df_pendencias.columns = df_pendencias.columns.str.strip()
        df_pendencias['CENTRO_CUSTO'] = df_pendencias['CENTRO_CUSTO'].astype(str).str.strip()
        if 'FILIAL' in df_pendencias.columns:
            df_pendencias['FILIAL'] = df_pendencias['FILIAL'].astype(str).str.strip().str.replace(r'\.0$', '', regex=True).str.zfill(6)
        if 'EMPRESA' in df_pendencias.columns:
            df_pendencias['EMPRESA'] = df_pendencias['EMPRESA'].astype(str).str.strip().str.replace(r'\.0$', '', regex=True).str.zfill(2)
    except Exception:
        df_pendencias = pd.DataFrame(columns=['EMPRESA', 'ORIGEM_DOC', 'FILIAL', 'NUMERO', 'EMISSAO', 'TOTAL_DOC', 'SITUACAO', 'CENTRO_CUSTO', 'DESC_CC', 'STATUS_BLOQ'])
    
    df_empresas = base[['COD EMPRESA Z01', 'DESC EMPRESA']].dropna(subset=['COD EMPRESA Z01'])
    df_empresas['COD EMPRESA Z01'] = df_empresas['COD EMPRESA Z01'].str.replace(r'\.0$', '', regex=True).str.zfill(2)
    
    df_filiais = base[['COD EMPRESA', 'COD FILIAL', 'DESC']].dropna(subset=['COD FILIAL'])
    df_filiais['COD EMPRESA'] = df_filiais['COD EMPRESA'].str.replace(r'\.0$', '', regex=True).str.zfill(2)
    df_filiais['COD FILIAL'] = df_filiais['COD FILIAL'].str.replace(r'\.0$', '', regex=True).str.zfill(6)
    
    plan2['CTT_EMPRESA'] = plan2['CTT_EMPRESA'].str.replace(r'\.0$', '', regex=True).str.zfill(2)
    plan2['CTT_FILIAL'] = plan2['CTT_FILIAL'].str.replace(r'\.0$', '', regex=True).str.zfill(6)
    plan2['CTT_CUSTO'] = plan2['CTT_CUSTO'].str.strip()
    
    base_unificada = df_filiais.merge(df_empresas, left_on='COD EMPRESA', right_on='COD EMPRESA Z01', how='left')
    return plan2, base_unificada, df_bloqueios, df_pendencias

if arquivo_unico:
    timestamp_atual = os.path.getmtime(arquivo_unico) if os.path.exists(arquivo_unico) and isinstance(arquivo_unico, str) else 0
    plan2, base_unificada, df_bloqueios, df_pendencias = carregar_dados(arquivo_unico, timestamp_atual)

    df_completo = plan2.merge(base_unificada, left_on=['CTT_EMPRESA', 'CTT_FILIAL'], right_on=['COD EMPRESA', 'COD FILIAL'], how='left')

    if not df_bloqueios.empty:
        df_completo = df_completo.merge(
            df_bloqueios,
            left_on=['CTT_EMPRESA', 'CTT_FILIAL', 'CTT_CUSTO'],
            right_on=['CTT_EMPRESA_BLOQ', 'CTT_FILIAL_BLOQ', 'CTT_CUSTO_BLOQ'],
            how='left'
        )
    else:
        df_completo['DATA_BLOQUEIO'] = pd.NA

    grp_cols = ['CENTRO_CUSTO']
    merge_left = ['CTT_CUSTO']
    if 'FILIAL' in df_pendencias.columns:
        grp_cols.append('FILIAL')
        merge_left.append('CTT_FILIAL')
    if 'EMPRESA' in df_pendencias.columns:
        grp_cols.append('EMPRESA')
        merge_left.append('CTT_EMPRESA')

    df_pendencias_count = df_pendencias.groupby(grp_cols).size().reset_index(name='PENDÊNCIAS_COUNT')
    df_completo = df_completo.merge(
        df_pendencias_count,
        left_on=merge_left,
        right_on=grp_cols,
        how='left'
    ).drop(columns=grp_cols, errors='ignore')
    df_completo['PENDÊNCIAS_COUNT'] = df_completo['PENDÊNCIAS_COUNT'].fillna(0).astype(int)

    # Força a remoção de espaços em todo o dataframe principal, ignorando qualquer dado cacheado anterior
    for col in df_completo.columns:
        if df_completo[col].dtype == 'object':
            df_completo[col] = df_completo[col].str.strip()

    df_completo['EMPRESA_COMPLETA'] = df_completo['COD EMPRESA Z01'].fillna('') + " - " + df_completo['DESC EMPRESA'].fillna('')
    df_completo['FILIAL_COMPLETA'] = df_completo['COD FILIAL'].fillna('') + " - " + df_completo['DESC'].fillna('')

    @st.dialog("Detalhes do Centro de Custo", width="large")
    def modal_detalhes_cc(ccusto_info, df_pendencias_geral):
        codigo_cc = ccusto_info['CTT_CUSTO']
        st.write(f"### Centro de Custo: {codigo_cc} - {ccusto_info['CTT_DESC01']}")
        st.write(f"**Filial:** {ccusto_info['FILIAL_COMPLETA']}")
        st.write("---")
        
        # Filtra o dataset gigante do UNION ALL apenas para o CC clicado
        mask = df_pendencias_geral['CENTRO_CUSTO'] == codigo_cc
        if 'FILIAL' in df_pendencias_geral.columns:
            mask &= (df_pendencias_geral['FILIAL'] == str(ccusto_info['CTT_FILIAL']))
        if 'EMPRESA' in df_pendencias_geral.columns:
            mask &= (df_pendencias_geral['EMPRESA'] == str(ccusto_info['CTT_EMPRESA']))
            
        df_pend_filtrado = df_pendencias_geral[mask]
        
        if not df_pend_filtrado.empty:
            st.warning(f"⚠️ Existem {len(df_pend_filtrado)} documento(s) em aberto para este Centro de Custo!")
            
            # Remove colunas desnecessárias (mantendo CENTRO_CUSTO no final para validação)
            df_display = df_pend_filtrado.drop(columns=['DESC_CC', 'STATUS_BLOQ', 'MOTIVO_INATIVO'], errors='ignore')
            
            # Garante que CENTRO_CUSTO seja a última coluna
            if 'CENTRO_CUSTO' in df_display.columns:
                cols = [c for c in df_display.columns if c != 'CENTRO_CUSTO'] + ['CENTRO_CUSTO']
                df_display = df_display[cols]
            
            # Formata data
            if 'EMISSAO' in df_display.columns:
                df_display['EMISSAO'] = pd.to_datetime(df_display['EMISSAO'], errors='coerce')
                
            # Formata valor monetário
            if 'TOTAL_DOC' in df_display.columns:
                df_display['TOTAL_DOC'] = pd.to_numeric(df_display['TOTAL_DOC'], errors='coerce').fillna(0)
            
            # Renomeia para o cabeçalho final
            df_display = df_display.rename(columns={
                'ORIGEM_DOC': 'ORIGEM DOC',
                'NUMERO': 'NUMERO DOC',
                'TOTAL_DOC': 'VALOR TOTAL',
                'CENTRO_CUSTO': 'CENTRO DE CUSTO'
            })
            
            # Garante a ordem exata solicitada
            ordem_desejada = ['ORIGEM DOC', 'FILIAL', 'EMISSAO', 'NUMERO DOC', 'VALOR TOTAL', 'SITUACAO', 'CENTRO DE CUSTO']
            colunas_presentes = [c for c in ordem_desejada if c in df_display.columns]
            extras = [c for c in df_display.columns if c not in colunas_presentes]
            df_display = df_display[colunas_presentes + extras]
                
            if 'VALOR TOTAL' in df_display.columns:
                df_show = df_display.style.format({
                    'VALOR TOTAL': lambda x: f"R$ {x:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.') if pd.notna(x) else ""
                })
            else:
                df_show = df_display

            st.dataframe(
                df_show, 
                hide_index=True, 
                use_container_width=True,
                column_config={
                    "EMISSAO": st.column_config.DatetimeColumn("EMISSAO", format="DD/MM/YYYY")
                }
            )
        else:
            st.success("✅ Nenhuma pendência encontrada para este Centro de Custo.")
            
        st.write("---")
        
        with st.expander("📅 Agendar Bloqueio", expanded=False):
            data_input = st.date_input("Data de bloqueio:", format="DD/MM/YYYY")
            
            st.write("**Autenticação Necessária**")
            usuario = st.text_input("Usuário")
            senha = st.text_input("Senha", type="password")
            
            st.markdown("<div style='margin-top: 15px;'></div>", unsafe_allow_html=True)
            salvar = st.button("Salvar", use_container_width=True)
            
            if salvar:
                admin_user = os.getenv("ADMIN_USER")
                admin_pwd = os.getenv("ADMIN_PWD")
                
                import sys
                sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..')))
                try:
                    import auth
                    db = auth.load_users_db()
                    user_clean = usuario.strip().lower()
                    tem_acesso = False
                    if usuario == admin_user and senha == admin_pwd:
                        tem_acesso = True
                    elif user_clean in db:
                        info = db[user_clean]
                        if not info.get("bloqueado") and info.get("senha") == auth.hash_password(senha):
                            if info.get("perfil") == "Administrador" or "Centro de Custo (Edição)" in info.get("paineis", []):
                                tem_acesso = True
                except Exception:
                    tem_acesso = (usuario == admin_user and senha == admin_pwd)

                if not tem_acesso:
                    st.error("Usuário ou senha incorretos, ou sem permissão de edição!")
                else:
                    try:
                        wb = openpyxl.load_workbook(caminho_fixo)
                        if 'Bloqueios' not in wb.sheetnames:
                            ws = wb.create_sheet('Bloqueios')
                            ws.append(['CTT_EMPRESA_BLOQ', 'CTT_FILIAL_BLOQ', 'CTT_CUSTO_BLOQ', 'DATA_BLOQUEIO'])
                        else:
                            ws = wb['Bloqueios']
                            if ws.cell(row=1, column=1).value != 'CTT_EMPRESA_BLOQ':
                                ws.delete_rows(1, ws.max_row)
                                ws.append(['CTT_EMPRESA_BLOQ', 'CTT_FILIAL_BLOQ', 'CTT_CUSTO_BLOQ', 'DATA_BLOQUEIO'])
                        
                        row_to_update = None
                        emp_alvo = str(ccusto_info['CTT_EMPRESA'])
                        fil_alvo = str(ccusto_info['CTT_FILIAL'])
                        cc_alvo_modal = str(ccusto_info['CTT_CUSTO'])
                        
                        for row in range(2, ws.max_row + 1):
                            if (str(ws.cell(row=row, column=1).value) == emp_alvo and 
                                str(ws.cell(row=row, column=2).value) == fil_alvo and
                                str(ws.cell(row=row, column=3).value) == cc_alvo_modal):
                                row_to_update = row
                                break
                        
                        data_formatada = data_input.strftime('%d/%m/%Y')
                        
                        if row_to_update:
                            ws.cell(row=row_to_update, column=4).value = data_formatada
                        else:
                            ws.append([emp_alvo, fil_alvo, cc_alvo_modal, data_formatada])
                            
                        wb.save(caminho_fixo)
                        st.success("Bloqueio agendado com sucesso!")
                        st.cache_data.clear()
                        st.rerun()
                    except PermissionError:
                        st.error("A planilha Aprovadores.xlsx está aberta. Feche-a para salvar.")
                    except Exception as e:
                        st.error("Erro interno ao salvar.")

    if 'filtro_empresa' not in st.session_state: st.session_state.filtro_empresa = 'TODAS'
    if 'filtro_filial' not in st.session_state: st.session_state.filtro_filial = 'TODAS'
    if 'filtro_regional' not in st.session_state: st.session_state.filtro_regional = 'TODAS'
    if 'filtro_busca' not in st.session_state: st.session_state.filtro_busca = ''
    if 'filtro_bloq' not in st.session_state: st.session_state.filtro_bloq = 'TODOS'
    if 'filtro_financeiro' not in st.session_state: st.session_state.filtro_financeiro = False
    if 'filtro_com_bloqueio' not in st.session_state: st.session_state.filtro_com_bloqueio = False
    if 'filtro_com_pendencia' not in st.session_state: st.session_state.filtro_com_pendencia = False

    def limpar_filtros():
        st.session_state.filtro_empresa = 'TODAS'
        st.session_state.filtro_filial = 'TODAS'
        st.session_state.filtro_regional = 'TODAS'
        st.session_state.filtro_busca = ''
        st.session_state.filtro_bloq = 'TODOS'
        st.session_state.filtro_financeiro = False
        st.session_state.filtro_com_bloqueio = False
        st.session_state.filtro_com_pendencia = False

    lista_empresas = sorted([e for e in df_completo['EMPRESA_COMPLETA'].dropna().unique().tolist() if e != " - "])
    lista_empresas.insert(0, 'TODAS')

    if st.session_state.filtro_empresa != 'TODAS':
        df_filiais_disponiveis = df_completo[df_completo['EMPRESA_COMPLETA'] == st.session_state.filtro_empresa]
    else:
        df_filiais_disponiveis = df_completo

    lista_filiais = sorted([f for f in df_filiais_disponiveis['FILIAL_COMPLETA'].dropna().unique().tolist() if f != " - "])
    lista_filiais.insert(0, 'TODAS')
    
    lista_regionais = sorted([str(r).strip() for r in df_completo['CTT_REGION'].dropna().unique().tolist() if str(r).strip() != ""])
    lista_regionais.insert(0, 'TODAS')
    
    opcoes_bloq = ['TODOS', 'ATIVO', 'INATIVO']

    import sys
    sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..')))
    try:
        import ui_utils
        ui_utils.render_standard_panel(
            title="Centro de Custo",
            subtitle="Relação de Centro de Custo",
            icon_name="C.Custo logo.png"
        )
    except ImportError:
        st.title("📋 Centro de Custo")
        st.markdown("---")
        st.markdown("")

    f_col1, f_col2, f_col3, f_col4, f_col5, f_col6 = st.columns([2, 2, 2, 2, 3, 1.5])
    with f_col1:
        st.selectbox("🏢 EMPRESA:", lista_empresas, key='filtro_empresa')
    with f_col2:
        st.selectbox("📍 FILIAL:", lista_filiais, key='filtro_filial')
    with f_col3:
        st.selectbox("🌎 REGIONAL:", lista_regionais, key='filtro_regional')
    with f_col4:
        st.selectbox("🟢 STATUS OPERACIONAL:", opcoes_bloq, key='filtro_bloq')
    with f_col5:
        st.text_input("🔍 Pesquisar Código ou Descrição...", key='filtro_busca')
    with f_col6:
        st.markdown('<div class="btn-limpar">', unsafe_allow_html=True)
        st.button("🗑️ Limpar Filtros", on_click=limpar_filtros, use_container_width=True)
        st.markdown('</div>', unsafe_allow_html=True)

    chk_col1, chk_col2 = st.columns(2)
    with chk_col1:
        st.checkbox("⚠️ Exibir apenas centros de custo com bloqueio agendado", key="filtro_com_bloqueio")
    with chk_col2:
        st.checkbox("🔴 Exibir apenas centros de custo com pendências", key="filtro_com_pendencia")

    st.write("")

    df_filtrado = df_completo.copy()
    if st.session_state.filtro_empresa != 'TODAS':
        df_filtrado = df_filtrado[df_filtrado['EMPRESA_COMPLETA'] == st.session_state.filtro_empresa]
    if st.session_state.filtro_filial != 'TODAS':
        df_filtrado = df_filtrado[df_filtrado['FILIAL_COMPLETA'] == st.session_state.filtro_filial]
    if st.session_state.filtro_regional != 'TODAS':
        df_filtrado = df_filtrado[df_filtrado['CTT_REGION'].astype(str).str.strip() == st.session_state.filtro_regional]
    if st.session_state.filtro_bloq != 'TODOS':
        df_filtrado = df_filtrado[df_filtrado['CTT_BLOQ'] == st.session_state.filtro_bloq]
    if st.session_state.get('filtro_financeiro', False):
        df_filtrado = df_filtrado[df_filtrado['CTT_XFINAN'].astype(str).str.strip().str.upper() == 'S']
    if st.session_state.get('filtro_com_bloqueio', False):
        df_filtrado = df_filtrado[
            (pd.notna(df_filtrado['DATA_BLOQUEIO'])) & 
            (df_filtrado['DATA_BLOQUEIO'].astype(str).str.strip() != '') &
            (df_filtrado['DATA_BLOQUEIO'].astype(str).str.strip().str.lower() != 'nan')
        ]
    if st.session_state.get('filtro_com_pendencia', False):
        df_filtrado = df_filtrado[df_filtrado['PENDÊNCIAS_COUNT'] > 0]
    if st.session_state.filtro_busca:
        def normalize(text):
            if not isinstance(text, str):
                text = str(text)
            # Remove acentos
            text = "".join(c for c in unicodedata.normalize('NFD', text) if unicodedata.category(c) != 'Mn')
            # Substitui hífens e sublinhados por espaços, remove espaços extras e converte para minúsculo
            return re.sub(r'[-_]', ' ', text).lower().strip()

        # Normaliza a busca completa e divide em palavras
        busca_full = normalize(st.session_state.filtro_busca)
        palavras_busca = busca_full.split()

        if palavras_busca:
            # Cria strings normalizadas para comparação (Código + Descrição)
            # Aplicar apply em cada linha pode ser lento em bases gigantes, mas para Dashboards costuma ser ok.
            # Uma otimização seria cachear a versão normalizada das colunas.
            norm_custo = df_filtrado['CTT_CUSTO'].apply(normalize)
            norm_desc = df_filtrado['CTT_DESC01'].apply(normalize)
            
            # Combina para buscar em ambos os campos ao mesmo tempo
            norm_combinado = norm_custo + " " + norm_desc
            
            # O registro bate se TODAS as palavras da busca forem encontradas no combinado
            mask = norm_combinado.apply(lambda x: all(p in x for p in palavras_busca))
            df_filtrado = df_filtrado[mask]

    total_cc = len(df_filtrado)
    ativos = len(df_filtrado[df_filtrado['CTT_BLOQ'] == 'ATIVO'])
    inativos = len(df_filtrado[df_filtrado['CTT_BLOQ'] == 'INATIVO'])
    financeiros = len(df_filtrado[df_filtrado['CTT_XFINAN'].astype(str).str.strip().str.upper() == 'S'])
    
    kpi_html = (
        '<div style="display: flex; gap: 14px; margin-bottom: 24px;">'

        '<div class="metric-card metric-card-blue" style="flex: 1; border-left-color: #3b82f6;">'
        '<div class="metric-info">'
        '<span class="metric-title">Total de Centros de Custo</span>'
        f'<span class="metric-value">{total_cc}</span>'
        '</div>'
        '<div class="metric-chart">'
        '<svg viewBox="0 0 100 30" width="60" height="30">'
        '<path d="M0,20 L20,25 L40,10 L60,15 L80,5 L100,0" fill="none" stroke="#3b82f6" stroke-width="2"/>'
        '</svg>'
        '</div>'
        '</div>'

        '<div class="metric-card metric-card-green" style="flex: 1; border-left-color: #10b981;">'
        '<div class="metric-info">'
        '<span class="metric-title">Centros de Custo Ativos</span>'
        f'<span class="metric-value" style="color: #10b981;">{ativos}</span>'
        '</div>'
        '<div class="metric-chart"><div class="donut-green"></div></div>'
        '</div>'

        '<div class="metric-card metric-card-red" style="flex: 1; border-left-color: #ef4444;">'
        '<div class="metric-info">'
        '<span class="metric-title">Centros de Custo Inativos</span>'
        f'<span class="metric-value" style="color: #ef4444;">{inativos}</span>'
        '</div>'
        '<div class="metric-chart"><div class="donut-red"></div></div>'
        '</div>'

        '<div class="metric-card metric-card-purple" style="flex: 1; border-left-color: #8b5cf6;">'
        '<div class="metric-info">'
        '<span class="metric-title">Somente Financeiro</span>'
        f'<span class="metric-value" style="color: #8b5cf6;">{financeiros}</span>'
        '</div>'
        '<div class="metric-chart"><div style="font-size: 28px;">💰</div></div>'
        '</div>'

        '</div>'
    )
    st.markdown(kpi_html, unsafe_allow_html=True)

    if df_filtrado.empty:
        st.warning("Nenhum registro encontrado para os filtros selecionados.")
    else:
        df_exibicao = df_filtrado.copy().fillna('')
        
        data_salvamento = datetime.fromtimestamp(timestamp_atual) if timestamp_atual > 0 else datetime.now()
        
        # Converte as colunas de data
        df_exibicao['DT INI EXIST'] = pd.to_datetime(df_exibicao['CTT_DTEXIS'], errors='coerce')
        df_exibicao['DT FIM EXIST'] = pd.to_datetime(df_exibicao['CTT_DTEXSF'], errors='coerce')
        
        data_salvamento_date = data_salvamento.date()
        
        def get_status_bolinha(row):
            dt_exis = row['DT INI EXIST']
            dt_exsf = row['DT FIM EXIST']
            status = row.get('CTT_BLOQ')
            
            if pd.notna(dt_exis) and dt_exis.date() > data_salvamento_date:
                return '🟡'
            if pd.notna(dt_exsf) and dt_exsf.date() < data_salvamento_date and status == 'ATIVO':
                return '⚪'
            if status == 'ATIVO':
                return '🟢'
            return '🔴'
        
        df_exibicao['LEGENDA'] = df_exibicao.apply(get_status_bolinha, axis=1)
        
        # Renomeando colunas para exibição na tabela nativa
        df_exibicao = df_exibicao.rename(columns={
            'EMPRESA_COMPLETA': 'EMPRESA',
            'FILIAL_COMPLETA': 'FILIAL',
            'CTT_CUSTO': 'CODIGO',
            'CTT_DESC01': 'DESCRIÇÃO',
            'CTT_BLOQ': 'STATUS',
            'CTT_XFINAN': 'FINANCEIRO?',
            'CTT_REGION': 'REGIONAL',
            'DATA_BLOQUEIO': 'DATA BLOQ',
            'PENDÊNCIAS_COUNT': 'PENDÊNCIAS'
        })
        
        colunas_exibir = ['LEGENDA', 'EMPRESA', 'FILIAL', 'CODIGO', 'DESCRIÇÃO', 'STATUS', 'FINANCEIRO?', 'REGIONAL', 'PENDÊNCIAS', 'DATA BLOQ', 'DT FIM EXIST']
        
        # Converte para datetime para permitir ordenação nativa (cronológica) pelo Streamlit
        df_exibicao['DATA BLOQ'] = pd.to_datetime(df_exibicao['DATA BLOQ'], format='%d/%m/%Y', errors='coerce')
        
        def destacar_linha(row):
            styles = [''] * len(row)

            if row['STATUS'] == 'ATIVO' and pd.notna(row['DATA BLOQ']):
                for i in range(len(row)):
                    styles[i] = 'background-color: rgba(245, 158, 11, 0.12); font-weight: 600'

            try:
                status_idx = colunas_exibir.index('STATUS')
                if row['STATUS'] == 'ATIVO':
                    styles[status_idx] += '; color: #10b981; font-weight: 700'
                elif row['STATUS'] == 'INATIVO':
                    styles[status_idx] += '; color: #ef4444; font-weight: 700'
            except ValueError:
                pass

            try:
                pend_idx = colunas_exibir.index('PENDÊNCIAS')
                if row['PENDÊNCIAS'] > 0:
                    styles[pend_idx] += '; color: #ef4444; font-weight: 700; background-color: rgba(239, 68, 68, 0.08);'
            except ValueError:
                pass

            return styles

        df_styled = (
            df_exibicao[colunas_exibir].style
            .apply(destacar_linha, axis=1)
            .format({
                'DATA BLOQ': lambda t: t.strftime('%d/%m/%Y') if pd.notna(t) else '-',
                'DT FIM EXIST': lambda t: t.strftime('%d/%m/%Y') if pd.notna(t) else '-'
            })
            .set_properties(subset=['LEGENDA'], **{'text-align': 'center'})
        )

        legend_html = """
        <div style="display: flex; justify-content: space-between; align-items: center; margin-bottom: 12px; flex-wrap: wrap; gap: 10px;">
            <div style="display: flex; align-items: center; gap: 18px; background: rgba(128,128,128,0.04); padding: 8px 18px; border-radius: 8px; border: 1px solid rgba(128,128,128,0.15); font-size: 12px; font-weight: 500; color: var(--text-color);">
                <span style="opacity: 0.5; font-size: 10px; font-weight: 700; text-transform: uppercase; letter-spacing: 0.5px; margin-right: -4px;">Legenda:</span>
                <div style="display: flex; align-items: center; gap: 6px;"><span style="font-size: 11px;">🟢</span> <span>Sem Restrição</span></div>
                <div style="display: flex; align-items: center; gap: 6px;"><span style="font-size: 11px;">🔴</span> <span>Bloqueado</span></div>
                <div style="display: flex; align-items: center; gap: 6px;"><span style="font-size: 11px;">🟡</span> <span>Exercício Não Iniciado</span></div>
                <div style="display: flex; align-items: center; gap: 6px;"><span style="font-size: 11px;">⚪</span> <span>Exercício Finalizado</span></div>
            </div>
            <div class="table-hint" style="margin-bottom: 0;">💡 Clique em qualquer linha para visualizar pendências e agendar bloqueios</div>
        </div>
        """
        st.markdown(legend_html, unsafe_allow_html=True)
        
        evento = st.dataframe(
            df_styled,
            hide_index=True,
            use_container_width=True,
            on_select="rerun",
            selection_mode="single-row",
            column_config={
                "LEGENDA": st.column_config.TextColumn(
                    "LEGENDA",
                    help="🟢 Sem Restrição\n\n🔴 Bloqueado\n\n🟡 Exercício Não Iniciado\n\n⚪ Exercício Finalizado",
                    width="small"
                ),
                "CODIGO": st.column_config.TextColumn(
                    "CODIGO",
                    width="small"
                ),
                "REGIONAL": st.column_config.TextColumn(
                    "REGIONAL",
                    width="medium"
                ),
                "PENDÊNCIAS": st.column_config.NumberColumn(
                    "PENDÊNCIAS",
                    help="Quantidade de documentos em aberto travando o processo de bloqueio"
                ),
                "DATA BLOQ": st.column_config.DateColumn(
                    "DATA BLOQ",
                    format="DD/MM/YYYY"
                ),
                "DT FIM EXIST": st.column_config.DateColumn(
                    "DT FIM EXIST",
                    format="DD/MM/YYYY"
                )
            }
        )
        
        linhas_selecionadas = evento.selection.rows
        if linhas_selecionadas:
            idx = linhas_selecionadas[0] # Índice (0-based) da linha selecionada
            ccusto_info = df_filtrado.iloc[idx]
            modal_detalhes_cc(ccusto_info, df_pendencias)

        st.write("")
        col_text, col_blank, col_btn = st.columns([4, 1, 2])
        
        with col_text:
            st.markdown(f'<div class="pag-text">Exibindo todos os <b>{len(df_exibicao)}</b> resultados encontrados.</div>', unsafe_allow_html=True)

        with col_btn:
            colunas_pra_remover = ['ST', 'CTT_EMPRESA', 'CTT_FILIAL', 'CTT_DESC01', 'CNPJ', 'CTT_EMPRESA Z01', 'CTT_DESC01', 'CTT_EMPRESA', 'CTT_FILIAL']
            df_export = df_exibicao.drop(columns=[col for col in colunas_pra_remover if col in df_exibicao.columns], errors='ignore')
            df_export.rename(columns={'EMPRESA_COMPLETA': 'EMPRESA', 'FILIAL_COMPLETA': 'FILIAL', 'CTT_CUSTO': 'C CUSTO', 'CTT_DESC01': 'DESCRICAO', 'CTT_BLOQ': 'BLOQUEADO?', 'CTT_XFINAN': 'FINANCEIRO?', 'CTT_REGION': 'REGIONAL'}, inplace=True)
            
            buffer = io.BytesIO()
            df_export.to_excel(buffer, index=False, engine='openpyxl')
            
            st.markdown('<div class="export-btn-container">', unsafe_allow_html=True)
            st.download_button(
                label="📊 Exportar Resultados (Excel)",
                data=buffer.getvalue(),
                file_name="filtro_centro_custo.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            st.markdown('</div>', unsafe_allow_html=True)
